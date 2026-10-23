"""Pasa la exportación de opiniones de producto de Trusted Shops (eTrusted, .xlsx) a un CSV para la tabla
store_product_reviews, SIN datos personales (ni nombre, ni correo, ni número de pedido).

    python3 scripts/trusted-reviews-to-csv.py privado/resenas-trusted-shops/resenas-producto-2026-10-23.xlsx privado/resenas-trusted-shops/resenas-producto.csv
    psql ... -c "\\copy public.store_product_reviews (review_id, reviewed_at, rating, title, body, reply, replied_at, status, gtin, old_sku, product_name, product_url) from 'privado/resenas-trusted-shops/resenas-producto.csv' csv header"
"""
import csv
import sys
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

COLUMNS = {
    "review_id": "ID de la valoración", "reviewed_at": "Creado a la/s*", "rating": "Valoración", "title": "Título de la valoración",
    "body": "Texto de la valoración", "reply": "Texto de respuesta", "replied_at": "Respondido a la/s*", "status": "Estado de la valoración",
    "gtin": "GTIN del producto", "old_sku": "SKU del producto", "product_name": "Nombre del producto", "product_url": "URL del producto",
}


def value(raw):
    if raw is None:
        return ""
    if isinstance(raw, datetime):
        return raw.isoformat()
    text = str(raw).strip()
    try:  # las fechas vienen como «31/01/26 23:46», en hora de Madrid
        return datetime.strptime(text, "%d/%m/%y %H:%M").replace(tzinfo=ZoneInfo("Europe/Madrid")).isoformat()
    except ValueError:
        pass
    return "" if text.lower() == "false" else text


rows = load_workbook(sys.argv[1], read_only=True).active.iter_rows(values_only=True)
header = next(row for row in rows if row and row[0] == "ID de la valoración")
index = {name: header.index(column) for name, column in COLUMNS.items()}
with open(sys.argv[2], "w", newline="", encoding="utf-8") as file:
    writer = csv.writer(file)
    writer.writerow(COLUMNS)
    count = 0
    for row in rows:
        if not row or not row[0] or row[index["rating"]] in (None, ""):
            continue
        record = {name: value(row[index[name]]) for name in COLUMNS}
        record["rating"] = str(round(float(record["rating"].replace(",", "."))))  # viene como «5.00»
        writer.writerow(record.values())
        count += 1
print(f"{count} opiniones")
